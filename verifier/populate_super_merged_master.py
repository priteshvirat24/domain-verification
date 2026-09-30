"""
Super Merged Master Domain Population & Discovery Pipeline:
1. Maps existing verified domains strictly by (Organization/Company Name, Country) from verified databases into SUPER_MERGED_MASTER_FINAL.xlsx. (Organization ID is intentionally NOT used due to duplicate/unreliable IDs).
2. Identifies remaining organizations lacking domain URLs.
3. Concurrently searches via Apify Google Search Scraper (28 parallel workers).
4. Concurrently crawls candidate domains via Scrapling (200 parallel workers).
5. Verifies domains with the Elimination-First Engine.
6. Updates and saves SUPER_MERGED_MASTER_FINAL.xlsx, SUPER_MERGED_MASTER_FINAL_POPULATED.xlsx, and exports final_super_merged_master_populated.csv.
"""
from __future__ import annotations

import asyncio
import csv
import json
import logging
import os
import re
import shutil
import sqlite3
import time
from collections import defaultdict
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl

from verifier.cache import SQLiteCache
from verifier.config import Config
from verifier.elimination_engine import evaluate_elimination_decision
from verifier.fetcher import TieredFetcher
from verifier.normalization import normalize_domain
from verifier.models import DomainRecord
from verifier.pipeline import investigate_domain
from verifier.apify_escalation import get_apify_client
from verifier.recover_inactive import (
    clean_org_for_search,
    extract_candidate_domain,
    run_batch_apify_search_with_retries,
    SearchCache,
    COUNTRY_MAP,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
LOG = logging.getLogger("super_merged_populator")


def load_known_domain_indexes() -> dict[tuple[str, str], str]:
    """Builds lookup dictionaries for existing verified domains strictly by (Name, Country)."""
    by_name_country: dict[tuple[str, str], str] = {}

    results_files = [
        "run_full_output/results.csv",
        "recovery_output/results_full_active_replaced.csv",
        "results_elimination.csv",
    ]

    for rf in results_files:
        p = Path(rf)
        if not p.exists():
            continue
        LOG.info("Indexing verified domains strictly by (Name, Country) from %s...", rf)
        with open(p, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for r in reader:
                domain = (r.get("Domain Name") or r.get("final_url") or "").strip()
                if not domain or domain.lower() in ("none", "nan", ""):
                    continue
                country = (r.get("Country") or r.get("\ufeffCountry") or "").strip().upper()
                for name_field in ("Organization Name", "organization_for_verification"):
                    val = (r.get(name_field) or "").strip()
                    if val and country:
                        by_name_country[(val.upper(), country)] = domain
                        by_name_country[(clean_org_for_search(val).upper(), country)] = domain

    LOG.info("Total strictly (name, country) keys indexed: %d", len(by_name_country))
    return by_name_country


async def main():
    excel_path = Path("SUPER_MERGED_MASTER_FINAL.xlsx")
    backup_path = Path("SUPER_MERGED_MASTER_FINAL_BACKUP.xlsx")

    if not backup_path.exists():
        LOG.info("Creating backup at %s...", backup_path)
        shutil.copyfile(excel_path, backup_path)

    # 1. Load existing domain indexes strictly by (Name, Country)
    by_name_country = load_known_domain_indexes()

    # 2. Connect to search cache and domain cache
    search_cache = SearchCache("verifier/search_cache.sqlite")
    cache_path = "verifier/live_validation_4000_cache.sqlite"
    config = Config(
        concurrency=200,
        max_evidence_pages=3,
        use_dynamic=False,
        use_stealth=False,
        use_proxy=False,
        use_apify=True,
    )
    cache = SQLiteCache(cache_path, config.cache_ttl_seconds)
    fetcher = TieredFetcher(config, proxy_url=None)
    apify_client = get_apify_client(config.apify_token)

    # 3. Load Excel Workbook
    LOG.info("Loading workbook %s...", excel_path)
    wb = openpyxl.load_workbook(excel_path)
    sheet_name = "final super merged master"
    ws = wb[sheet_name]
    LOG.info("Workbook loaded. Rows: %d", ws.max_row)

    # Column mappings for 'final super merged master'
    col_domain = 13
    col_org_name = 6
    col_orig_name = 7
    col_country = 2
    col_country_group = 1

    already_had_domain = 0
    mapped_from_existing = 0
    missing_rows = []

    for row_idx in range(2, ws.max_row + 1):
        cell_domain = ws.cell(row=row_idx, column=col_domain).value
        curr_domain = str(cell_domain or "").strip()
        if curr_domain and curr_domain.lower() not in ("none", "nan", ""):
            already_had_domain += 1
            continue

        org_name = str(ws.cell(row=row_idx, column=col_org_name).value or "").strip()
        orig_name = str(ws.cell(row=row_idx, column=col_orig_name).value or "").strip()
        country_code = str(ws.cell(row=row_idx, column=col_country).value or ws.cell(row=row_idx, column=col_country_group).value or "").strip().upper()

        found_domain = None
        for candidate_name in (org_name, orig_name):
            if not candidate_name or candidate_name.lower() in ("none", "nan", ""):
                continue
            c_upper = candidate_name.upper()
            c_clean = clean_org_for_search(candidate_name).upper()
            if (c_upper, country_code) in by_name_country:
                found_domain = by_name_country[(c_upper, country_code)]
                break
            elif (c_clean, country_code) in by_name_country:
                found_domain = by_name_country[(c_clean, country_code)]
                break

        if found_domain:
            ws.cell(row=row_idx, column=col_domain, value=found_domain)
            mapped_from_existing += 1
        else:
            primary_name = org_name or orig_name
            if primary_name and primary_name.lower() not in ("none", "nan", ""):
                missing_rows.append({
                    "row_idx": row_idx,
                    "org_name": primary_name,
                    "country_code": country_code,
                })

    LOG.info("Phase 1 Strict Name+Country Mapping Complete:")
    LOG.info("  Already had domain: %d", already_had_domain)
    LOG.info("  Newly mapped strictly by (Name, Country): %d", mapped_from_existing)
    LOG.info("  Total now populated: %d / %d (%.2f%%)",
             already_had_domain + mapped_from_existing, ws.max_row - 1,
             ((already_had_domain + mapped_from_existing) / (ws.max_row - 1)) * 100)
    LOG.info("  Remaining rows needing Apify/Scrapling discovery: %d", len(missing_rows))

    # Save interim mapping update
    wb.save(excel_path)
    LOG.info("Saved interim strictly mapped progress to %s", excel_path)

    # 4. Phase 2: Launch Apify & Scrapling Discovery for remaining missing rows
    if not missing_rows:
        LOG.info("All rows have domain URLs populated! Pipeline complete.")
        return

    # Deduplicate search queries
    query_to_missing_rows = defaultdict(list)
    for r in missing_rows:
        country_name = COUNTRY_MAP.get(r["country_code"], r["country_code"])
        clean_org = clean_org_for_search(r["org_name"])
        q = f'"{clean_org}" {country_name} official website'.strip()
        query_to_missing_rows[q].append(r)

    all_queries = list(query_to_missing_rows.keys())
    needed_queries = []
    search_results_by_query = {}

    for q in all_queries:
        cached_res = search_cache.get(q)
        if cached_res is not None:
            search_results_by_query[q] = cached_res
        else:
            needed_queries.append(q)

    LOG.info("Search Cache Status: %d queries cached, %d queries need fetching via Apify",
             len(search_results_by_query), len(needed_queries))

    # Run missing queries via Apify in batches of 50 with 28 concurrent workers
    batch_size = 50
    sem_apify = asyncio.Semaphore(28)

    async def fetch_chunk(chunk: list[str], chunk_idx: int, total_chunks: int):
        async with sem_apify:
            LOG.info("Apify Search chunk %d/%d (%d queries)...", chunk_idx, total_chunks, len(chunk))
            batch_res = await asyncio.to_thread(run_batch_apify_search_with_retries, apify_client, chunk)
            for q, org_res in batch_res.items():
                search_cache.put(q, org_res)
                search_results_by_query[q] = org_res
            for q in chunk:
                if q not in search_results_by_query:
                    search_cache.put(q, [])
                    search_results_by_query[q] = []

    chunks = [needed_queries[i:i + batch_size] for i in range(0, len(needed_queries), batch_size)]
    if chunks:
        LOG.info("Dispatching %d search chunks to 28 parallel Apify workers...", len(chunks))
        await asyncio.gather(*(fetch_chunk(chunk, idx + 1, len(chunks)) for idx, chunk in enumerate(chunks)))

    # 5. Extract candidate domains and crawl
    LOG.info("Extracting candidate domains from SERP results...")
    candidate_domains_by_query = {}
    unique_candidates = set()

    for q in all_queries:
        organic = search_results_by_query.get(q, [])
        cand = extract_candidate_domain(organic)
        if cand:
            candidate_domains_by_query[q] = cand
            unique_candidates.add(cand)

    LOG.info("Candidate domains discovered: %d unique hosts", len(unique_candidates))

    # Concurrently crawl candidate domains with Scrapling (200 workers)
    sem_crawl = asyncio.Semaphore(config.concurrency)

    async def crawl_domain(dom: str):
        async with sem_crawl:
            try:
                norm = normalize_domain(dom)
                host = norm.normalized_domain
                if not host:
                    return
                rec = await cache.get(host)
                if not rec:
                    rec = await investigate_domain(host, config, fetcher)
                    await cache.put(rec)
            except Exception as e:
                LOG.warning("Error crawling candidate domain %s: %s", dom, e)

    LOG.info("Crawling unique candidate domains with 200 workers...")
    await asyncio.gather(*(crawl_domain(d) for d in unique_candidates), return_exceptions=True)

    # 6. Verify and populate candidate domains into Excel workbook
    LOG.info("Verifying candidate domains and populating workbook...")
    newly_verified_count = 0

    for q, rows_list in query_to_missing_rows.items():
        cand = candidate_domains_by_query.get(q)
        if not cand:
            continue

        norm = normalize_domain(cand)
        host = norm.normalized_domain
        dom_rec = await cache.get(host)

        # Evaluate verification decision
        decision = evaluate_elimination_decision(
            {"Organization Name": rows_list[0]["org_name"], "Country": rows_list[0]["country_code"]},
            norm,
            dom_rec,
            network_healthy=True,
        )

        decision_class = decision.get("classification")
        if decision_class in ("VALID", "VALID_GROUP"):
            domain_url = f"https://{host}"
            for r in rows_list:
                ws.cell(row=r["row_idx"], column=col_domain, value=domain_url)
                newly_verified_count += 1

    # Save final updated workbook
    LOG.info("Saving final populated workbook %s...", excel_path)
    wb.save(excel_path)
    wb.save("SUPER_MERGED_MASTER_FINAL_POPULATED.xlsx")

    # Export to CSV for fast downstream consumption
    csv_out = Path("final_super_merged_master_populated.csv")
    with open(csv_out, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for r in ws.iter_rows(values_only=True):
            writer.writerow(r)

    total_final_populated = already_had_domain + mapped_from_existing + newly_verified_count
    total_data_rows = ws.max_row - 1
    final_rate = (total_final_populated / total_data_rows) * 100

    LOG.info("==================================================")
    LOG.info("SUPER MERGED MASTER POPULATION COMPLETE")
    LOG.info("Total Data Rows: %d", total_data_rows)
    LOG.info("Previously Populated: %d", already_had_domain)
    LOG.info("Strictly Mapped by (Name, Country): %d", mapped_from_existing)
    LOG.info("Newly Discovered via Apify + Scrapling: %d", newly_verified_count)
    LOG.info("Total Final Populated: %d (%.2f%%)", total_final_populated, final_rate)
    LOG.info("Saved to: %s and %s", excel_path, csv_out)
    LOG.info("==================================================")


if __name__ == "__main__":
    asyncio.run(main())
