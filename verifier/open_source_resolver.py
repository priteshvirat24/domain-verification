"""
Enterprise Open-Source Domain Discovery & Recovery Engine:
1. Completely replaces Apify with Camoufox (C++ Patched Anti-Detect Firefox Engine).
2. Stealthily harvests organic SERPs via Bing without CAPTCHAs, API costs, or monthly quotas.
3. Automatically decodes redirect URLs and extracts verified corporate homepages.
4. Crawls candidate domains via Scrapling (curl_cffi TLS/HTTP2 browser impersonation, 150 workers).
5. Validates candidate domains with the 9-point Elimination-First Engine.
6. Populates and updates SUPER_MERGED_MASTER_FINAL.xlsx and final_super_merged_master_populated.csv.
"""
from __future__ import annotations

import asyncio
import base64
import csv
import json
import logging
import os
import random
import re
import shutil
import sqlite3
import time
from pathlib import Path
from urllib.parse import quote_plus, urlsplit

import openpyxl
from camoufox.async_api import AsyncCamoufox

from verifier.cache import SQLiteCache
from verifier.config import Config
from verifier.elimination_engine import evaluate_elimination_decision
from verifier.fetcher import TieredFetcher
from verifier.models import DomainRecord
from verifier.normalization import normalize_domain
from verifier.pipeline import investigate_domain
from verifier.recover_inactive import (
    clean_org_for_search,
    extract_candidate_domain,
    SearchCache,
    COUNTRY_MAP,
    EXCLUDE_DOMAINS,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
LOG = logging.getLogger("camoufox_resolver")


def decode_bing_url(href: str) -> str:
    """Decodes Bing redirect links (&u=a1...) to retrieve canonical target URL."""
    if "/ck/a?!" in href and "&u=" in href:
        try:
            part = href.split("&u=")[1].split("&")[0]
            if part.startswith("a1"):
                raw = part[2:] + "==="
                raw = raw[:len(raw) - (len(raw) % 4)] if len(raw) % 4 != 0 else raw
                return base64.urlsafe_b64decode(raw).decode("utf-8", errors="ignore")
        except Exception:
            pass
    return href


async def search_camoufox_batch(queries: list[str], search_cache: SearchCache, concurrency: int = 6) -> dict[str, list[dict]]:
    """Runs a batch of search queries concurrently using a shared Camoufox anti-detect browser."""
    results_map = {}
    queries_to_fetch = [q for q in queries if search_cache.get(q) is None]
    
    LOG.info("Total queries: %d (Already cached: %d, To fetch: %d)",
             len(queries), len(queries) - len(queries_to_fetch), len(queries_to_fetch))
    
    for q in queries:
        cached = search_cache.get(q)
        if cached is not None:
            results_map[q] = cached

    if not queries_to_fetch:
        return results_map

    sem = asyncio.Semaphore(concurrency)
    completed = 0

    async with AsyncCamoufox(headless=True) as browser:
        context = await browser.new_context()

        async def fetch_one(q: str):
            nonlocal completed
            async with sem:
                page = await context.new_page()
                try:
                    url = f"https://www.bing.com/search?q={quote_plus(q)}"
                    await page.goto(url, wait_until="domcontentloaded", timeout=15000)
                    links = await page.locator("li.b_algo h2 a").all()
                    org_results = []
                    for l in links[:5]:
                        href = await l.get_attribute("href") or ""
                        title = await l.inner_text()
                        real_url = decode_bing_url(href)
                        if real_url and real_url.startswith("http"):
                            org_results.append({"title": title, "url": real_url})
                    search_cache.put(q, org_results)
                    results_map[q] = org_results
                except Exception as e:
                    LOG.debug("Camoufox query error for '%s': %s", q, e)
                    search_cache.put(q, [])
                    results_map[q] = []
                finally:
                    await page.close()
                    completed += 1
                    if completed % 100 == 0 or completed == len(queries_to_fetch):
                        LOG.info("Camoufox Search Progress: %d / %d queries completed...", completed, len(queries_to_fetch))

        await asyncio.gather(*(fetch_one(q) for q in queries_to_fetch))

    return results_map


async def main():
    excel_path = Path("SUPER_MERGED_MASTER_FINAL.xlsx")
    csv_path = Path("final_super_merged_master_populated.csv")

    LOG.info("1. Loading unresolved entities from %s...", csv_path)
    all_rows = []
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        all_rows = list(reader)

    unresolved_rows = []
    for idx, r in enumerate(all_rows, 2):
        r["_row_idx"] = idx
        st = r.get("Domain Status")
        if st in ("UNVERIFIED_INACTIVE", "NOT_FOUND"):
            unresolved_rows.append(r)

    LOG.info("Targeting %d unresolved rows with Camoufox + Scrapling", len(unresolved_rows))

    search_cache = SearchCache("verifier/search_cache.sqlite")
    config = Config(concurrency=150, max_evidence_pages=2, use_dynamic=False, use_stealth=False, use_proxy=False, use_apify=False)
    domain_cache = SQLiteCache("verifier/live_validation_4000_cache.sqlite", config.cache_ttl_seconds)
    fetcher = TieredFetcher(config)

    # Build queries
    query_to_rows = {}
    for r in unresolved_rows:
        name = (r.get("Organization Name") or r.get("original_company_name") or "").strip()
        if not name or name.lower() in ("none", "nan", ""):
            continue
        c_code = (r.get("country code") or r.get("Country (Group)") or "").strip().upper()
        country_name = COUNTRY_MAP.get(c_code, c_code)
        clean = clean_org_for_search(name)
        tax_id = (r.get("tax_id") or "").strip()

        if tax_id and len(tax_id) >= 6:
            q = f'"{clean}" {tax_id} official website'
        else:
            q = f'{clean} {country_name} official website'

        if q not in query_to_rows:
            query_to_rows[q] = []
        query_to_rows[q].append(r)

    unique_queries = list(query_to_rows.keys())
    LOG.info("2. Harvesting organic SERPs for %d unique queries via Camoufox...", len(unique_queries))
    search_results = await search_camoufox_batch(unique_queries, search_cache, concurrency=8)

    # Extract candidates
    LOG.info("3. Extracting candidate domains from harvested SERPs...")
    candidate_domains_by_query = {}
    unique_candidates = set()

    for q, org_res in search_results.items():
        cand = extract_candidate_domain(org_res)
        if cand:
            candidate_domains_by_query[q] = cand
            unique_candidates.add(cand)

    LOG.info("Discovered %d unique candidate domains across %d queries",
             len(unique_candidates), len(candidate_domains_by_query))

    # Crawl candidates with Scrapling (150 workers)
    LOG.info("4. Crawling %d unique candidate hosts with 150 Scrapling workers...", len(unique_candidates))
    sem_crawl = asyncio.Semaphore(150)

    async def crawl_domain(dom: str):
        async with sem_crawl:
            try:
                norm = normalize_domain(dom)
                host = norm.normalized_domain
                if not host:
                    return
                rec = await domain_cache.get(host)
                if not rec:
                    rec = await investigate_domain(host, config, fetcher)
                    await domain_cache.put(rec)
            except Exception as e:
                LOG.debug("Crawl error for %s: %s", dom, e)

    await asyncio.gather(*(crawl_domain(d) for d in unique_candidates), return_exceptions=True)
    LOG.info("Crawl completed!")

    # Evaluate Elimination decisions
    LOG.info("5. Evaluating decisions and updating records...")
    newly_replaced = 0
    newly_populated = 0

    for q, rows_list in query_to_rows.items():
        cand = candidate_domains_by_query.get(q)
        if not cand:
            continue
        norm = normalize_domain(cand)
        host = norm.normalized_domain
        dom_rec = await domain_cache.get(host)
        if not dom_rec:
            continue

        decision = evaluate_elimination_decision(
            {"Organization Name": rows_list[0].get("Organization Name") or rows_list[0].get("original_company_name"),
             "Country": rows_list[0].get("country code")},
            norm,
            dom_rec,
            network_healthy=True,
        )

        decision_class = decision.get("classification")
        if decision_class in ("VALID", "VALID_GROUP"):
            domain_url = f"https://{host}"
            for r in rows_list:
                prior_st = r.get("Domain Status")
                r["Domain URL"] = domain_url
                if prior_st == "UNVERIFIED_INACTIVE":
                    r["Domain Status"] = "REPLACED_INACTIVE"
                    newly_replaced += 1
                else:
                    r["Domain Status"] = "VERIFIED_ACTIVE"
                    newly_populated += 1

    LOG.info("Evaluation Complete: %d Inactive Replaced, %d Empty Populated!", newly_replaced, newly_populated)

    # Save to Excel & CSV
    LOG.info("6. Saving updated master workbook to %s...", excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["final super merged master"]

    col_domain = 13
    col_legacy = 14
    col_status = 15

    total_active = 0
    status_summary = {}

    for r in all_rows:
        row_idx = r["_row_idx"]
        d = r.get("Domain URL")
        st = r.get("Domain Status")
        leg = r.get("Original Legacy Domain")

        ws.cell(row=row_idx, column=col_domain, value=d)
        ws.cell(row=row_idx, column=col_legacy, value=leg)
        ws.cell(row=row_idx, column=col_status, value=st)

        status_summary[st] = status_summary.get(st, 0) + 1
        if st in ("VERIFIED_ACTIVE", "REPLACED_INACTIVE", "PRE_EXISTING_ACTIVE"):
            total_active += 1

    wb.save(excel_path)
    wb.save("SUPER_MERGED_MASTER_FINAL_POPULATED.xlsx")

    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        fieldnames = [c for c in all_rows[0].keys() if c != "_row_idx"]
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(all_rows)

    total_rows = len(all_rows)
    coverage = round((total_active / total_rows) * 100, 2)

    summary = {
        "total_rows": total_rows,
        "newly_replaced_inactive": newly_replaced,
        "newly_populated_empty": newly_populated,
        "total_active_domains_now": total_active,
        "active_domain_coverage_percent": coverage,
        "status_breakdown": status_summary,
    }

    with open("camoufox_recovery_summary.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2)

    LOG.info("==================================================")
    LOG.info("CAMOUFOX OPEN-SOURCE RESOLUTION ENGINE FINISHED")
    LOG.info("Newly Replaced Inactive: %d", newly_replaced)
    LOG.info("Newly Populated Empty: %d", newly_populated)
    LOG.info("TOTAL ACTIVE DOMAINS NOW: %d (%.2f%%)", total_active, coverage)
    LOG.info("Status Breakdown: %s", status_summary)
    LOG.info("Saved to: %s and %s", excel_path, csv_path)
    LOG.info("==================================================")


if __name__ == "__main__":
    asyncio.run(main())
